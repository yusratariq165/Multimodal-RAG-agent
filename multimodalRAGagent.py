import os
import tempfile
import base64
from pathlib import Path

import streamlit as st
import pandas as pd
import docx
from pptx import Presentation

from openai import OpenAI

from langchain_core.documents import Document
from langchain_core.output_parsers import StrOutputParser
from langchain_core.prompts import ChatPromptTemplate
from langchain_core.runnables import RunnablePassthrough

from langchain_openai import ChatOpenAI, OpenAIEmbeddings

from langchain_community.document_loaders import (
    PyPDFLoader,
    TextLoader
)

from langchain_community.vectorstores import FAISS
from langchain_text_splitters import RecursiveCharacterTextSplitter


# =========================================================
# PAGE CONFIGURATION
# =========================================================

st.set_page_config(
    page_title="Multiformat RAG Assistant",
    page_icon="🤖",
    layout="wide"
)


# =========================================================
# SESSION STATE
# =========================================================

if "api_key_verified" not in st.session_state:
    st.session_state.api_key_verified = False

if "api_key" not in st.session_state:
    st.session_state.api_key = ""

if "vectorstore" not in st.session_state:
    st.session_state.vectorstore = None

if "messages" not in st.session_state:
    st.session_state.messages = []

if "file_sig" not in st.session_state:
    st.session_state.file_sig = None


# =========================================================
# API KEY GATE
# =========================================================

if not st.session_state.api_key_verified:

    st.title("🔐 Multiformat RAG Assistant")

    st.subheader("OpenAI API Key Required")

    st.write(
        "Enter your OpenAI API key to access the RAG Assistant."
    )

    st.info(
        "Your API key is required before you can upload "
        "and process files."
    )

    api_key_input = st.text_input(
        "🔑 Enter OpenAI API Key",
        type="password",
        placeholder="sk-..."
    )

    if st.button(
        "🔓 Verify & Continue",
        use_container_width=True
    ):

        if not api_key_input.strip():

            st.error(
                "❌ Please enter your API key."
            )

        else:

            try:

                # -----------------------------------------
                # Store API key temporarily
                # -----------------------------------------

                clean_key = api_key_input.strip()

                os.environ["OPENAI_API_KEY"] = clean_key


                # -----------------------------------------
                # Test API key
                # -----------------------------------------

                test_client = OpenAI(
                    api_key=clean_key
                )

                test_client.models.list()


                # -----------------------------------------
                # API key is valid
                # -----------------------------------------

                st.session_state.api_key = clean_key

                st.session_state.api_key_verified = True

                st.success(
                    "✅ API key verified successfully!"
                )

                st.rerun()


            except Exception as e:

                st.error(
                    "❌ API key verification failed."
                )

                st.warning(
                    f"Details: {str(e)}"
                )


    st.markdown("---")

    st.caption(
        "🔐 API Key Gate | Multiformat RAG Assistant"
    )

    # Stop application here.
    # Nothing below will run until API key is verified.

    st.stop()


# =========================================================
# API KEY RESTORE
# =========================================================

os.environ["OPENAI_API_KEY"] = (
    st.session_state.api_key
)


# =========================================================
# MAIN TITLE
# =========================================================

st.title("🤖 Multiformat RAG Assistant")

st.write(
    "Upload a document, image, or audio file "
    "and ask questions about it."
)


# =========================================================
# SIDEBAR
# =========================================================

st.sidebar.header("⚙️ Settings")


# =========================================================
# API KEY STATUS
# =========================================================

st.sidebar.success(
    "🟢 API Key Verified"
)


if st.sidebar.button(
    "🔑 Change API Key",
    use_container_width=True
):

    st.session_state.api_key_verified = False
    st.session_state.api_key = ""
    st.session_state.vectorstore = None
    st.session_state.messages = []
    st.session_state.file_sig = None

    os.environ.pop(
        "OPENAI_API_KEY",
        None
    )

    st.rerun()


# =========================================================
# MODEL SELECTION
# =========================================================

model_name = st.sidebar.selectbox(
    "🤖 Select Model",
    [
        "gpt-4o-mini",
        "gpt-4o",
        "gpt-4.1-mini",
        "gpt-3.5-turbo"
    ]
)


# =========================================================
# CHUNK SIZE
# =========================================================

chunk_size = st.sidebar.slider(
    "📦 Chunk Size",
    min_value=500,
    max_value=2000,
    value=1000,
    step=100
)


# =========================================================
# CHUNK OVERLAP
# =========================================================

chunk_overlap = st.sidebar.slider(
    "🔗 Chunk Overlap",
    min_value=0,
    max_value=400,
    value=150,
    step=50
)


# =========================================================
# TOP K
# =========================================================

top_k = st.sidebar.slider(
    "🔎 Retrieved Chunks",
    min_value=2,
    max_value=10,
    value=4
)


# =========================================================
# CLEAR CHAT
# =========================================================

if st.sidebar.button(
    "🗑️ Clear Chat",
    use_container_width=True
):

    st.session_state.messages = []

    st.rerun()


# =========================================================
# SUPPORTED FILE FORMATS
# =========================================================

SUPPORTED_FILES = [
    "pdf",
    "docx",
    "xlsx",
    "xls",
    "csv",
    "txt",
    "pptx",
    "jpg",
    "jpeg",
    "png",
    "mp3",
    "wav"
]


# =========================================================
# PROCESS PDF
# =========================================================

def process_pdf(file_path):

    loader = PyPDFLoader(file_path)

    documents = loader.load()

    for doc in documents:

        doc.metadata["file_type"] = "PDF"

    return documents


# =========================================================
# PROCESS DOCX
# =========================================================

def process_docx(file_path):

    document = docx.Document(file_path)

    text = []

    for paragraph in document.paragraphs:

        if paragraph.text.strip():

            text.append(
                paragraph.text
            )

    full_text = "\n".join(text)

    return [
        Document(
            page_content=full_text,
            metadata={
                "source": file_path,
                "file_type": "DOCX"
            }
        )
    ]


# =========================================================
# PROCESS EXCEL
# =========================================================

def process_excel(file_path):

    df = pd.read_excel(file_path)

    text = df.to_string(
        index=False
    )

    return [
        Document(
            page_content=text,
            metadata={
                "source": file_path,
                "file_type": "Excel"
            }
        )
    ]


# =========================================================
# PROCESS CSV
# =========================================================

def process_csv(file_path):

    df = pd.read_csv(file_path)

    text = df.to_string(
        index=False
    )

    return [
        Document(
            page_content=text,
            metadata={
                "source": file_path,
                "file_type": "CSV"
            }
        )
    ]


# =========================================================
# PROCESS TXT
# =========================================================

def process_txt(file_path):

    loader = TextLoader(
        file_path,
        encoding="utf-8"
    )

    documents = loader.load()

    for doc in documents:

        doc.metadata["file_type"] = "TXT"

    return documents


# =========================================================
# PROCESS POWERPOINT
# =========================================================

def process_pptx(file_path):

    presentation = Presentation(
        file_path
    )

    slides_text = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        slide_text = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                if shape.text.strip():

                    slide_text.append(
                        shape.text
                    )

        if slide_text:

            slides_text.append(
                f"Slide {slide_number}:\n"
                + "\n".join(slide_text)
            )


    full_text = "\n\n".join(
        slides_text
    )


    return [
        Document(
            page_content=full_text,
            metadata={
                "source": file_path,
                "file_type": "PPTX"
            }
        )
    ]


# =========================================================
# PROCESS IMAGE
# =========================================================

def process_image(
    file_path,
    model_name
):

    # -----------------------------------------
    # Read image
    # -----------------------------------------

    with open(
        file_path,
        "rb"
    ) as image_file:

        image_data = base64.b64encode(
            image_file.read()
        ).decode("utf-8")


    # -----------------------------------------
    # Detect image type
    # -----------------------------------------

    extension = Path(
        file_path
    ).suffix.lower()


    if extension in [
        ".jpg",
        ".jpeg"
    ]:

        mime_type = "image/jpeg"

    else:

        mime_type = "image/png"


    # -----------------------------------------
    # Vision model
    # -----------------------------------------

    vision_model = ChatOpenAI(
        model=model_name,
        temperature=0
    )


    # -----------------------------------------
    # Multimodal message
    # -----------------------------------------

    message = [
        {
            "type": "text",
            "text": """
Analyze this image carefully.

Extract all useful information.

If there is text, read it.

If there is a table, describe
its contents.

If there is a chart, explain
the important information.

If there is a diagram, explain
its components and relationships.

Return a detailed description
that can be used for question
answering later.
"""
        },
        {
            "type": "image_url",
            "image_url": {
                "url":
                    f"data:{mime_type};base64,{image_data}"
            }
        }
    ]


    # -----------------------------------------
    # Send image to model
    # -----------------------------------------

    response = vision_model.invoke(
        [
            {
                "role": "user",
                "content": message
            }
        ]
    )


    return [
        Document(
            page_content=response.content,
            metadata={
                "source": file_path,
                "file_type": "Image"
            }
        )
    ]


# =========================================================
# PROCESS AUDIO
# =========================================================

def process_audio(file_path):

    client = OpenAI(
        api_key=os.environ[
            "OPENAI_API_KEY"
        ]
    )


    with open(
        file_path,
        "rb"
    ) as audio_file:

        transcript = (
            client.audio.transcriptions.create(
                model="gpt-4o-mini-transcribe",
                file=audio_file
            )
        )


    return [
        Document(
            page_content=transcript.text,
            metadata={
                "source": file_path,
                "file_type": "Audio"
            }
        )
    ]


# =========================================================
# LOAD FILE
# =========================================================

def load_file(
    file_path,
    extension,
    model_name
):

    extension = extension.lower()


    if extension == "pdf":

        return process_pdf(
            file_path
        )


    elif extension == "docx":

        return process_docx(
            file_path
        )


    elif extension in [
        "xlsx",
        "xls"
    ]:

        return process_excel(
            file_path
        )


    elif extension == "csv":

        return process_csv(
            file_path
        )


    elif extension == "txt":

        return process_txt(
            file_path
        )


    elif extension == "pptx":

        return process_pptx(
            file_path
        )


    elif extension in [
        "jpg",
        "jpeg",
        "png"
    ]:

        return process_image(
            file_path,
            model_name
        )


    elif extension in [
        "mp3",
        "wav"
    ]:

        return process_audio(
            file_path
        )


    else:

        raise ValueError(
            f"Unsupported file format: .{extension}"
        )


# =========================================================
# BUILD VECTOR STORE
# =========================================================

def build_vectorstore(
    uploaded_file,
    model_name,
    chunk_size,
    chunk_overlap
):

    extension = Path(
        uploaded_file.name
    ).suffix.lower().replace(
        ".",
        ""
    )


    # -----------------------------------------
    # Temporary file
    # -----------------------------------------

    with tempfile.NamedTemporaryFile(
        delete=False,
        suffix=f".{extension}"
    ) as temp_file:

        temp_file.write(
            uploaded_file.getbuffer()
        )

        temp_path = temp_file.name


    try:

        # -------------------------------------
        # Extract content
        # -------------------------------------

        documents = load_file(
            temp_path,
            extension,
            model_name
        )


        if not documents:

            raise ValueError(
                "No readable content found."
            )


        # -------------------------------------
        # Original filename
        # -------------------------------------

        for document in documents:

            document.metadata[
                "original_file"
            ] = uploaded_file.name


        # -------------------------------------
        # Split into chunks
        # -------------------------------------

        splitter = (
            RecursiveCharacterTextSplitter(
                chunk_size=chunk_size,
                chunk_overlap=chunk_overlap
            )
        )


        chunks = splitter.split_documents(
            documents
        )


        if not chunks:

            raise ValueError(
                "Could not create document chunks."
            )


        # -------------------------------------
        # Create embeddings
        # -------------------------------------

        embeddings = OpenAIEmbeddings(
            model="text-embedding-3-small"
        )


        # -------------------------------------
        # FAISS vector database
        # -------------------------------------

        vectorstore = FAISS.from_documents(
            chunks,
            embeddings
        )


        return vectorstore


    finally:

        # -------------------------------------
        # Delete temporary file
        # -------------------------------------

        try:

            os.remove(
                temp_path
            )

        except Exception:

            pass


# =========================================================
# RAG PROMPT
# =========================================================

RAG_PROMPT = ChatPromptTemplate.from_template(
    """
You are a document question-answering assistant.

Answer the user's question ONLY using
the provided context.

Do NOT use outside knowledge.

If the answer cannot be found in the
provided context, respond exactly:

I couldn't find that in the document.

Keep the answer clear and easy to understand.

If source information is available,
mention the source or page number.

==============================
CONTEXT
==============================

{context}

==============================
QUESTION
==============================

{question}
"""
)


# =========================================================
# FORMAT DOCUMENTS
# =========================================================

def format_docs(documents):

    formatted = []


    for document in documents:

        source = document.metadata.get(
            "original_file",
            document.metadata.get(
                "source",
                "Unknown"
            )
        )


        file_type = document.metadata.get(
            "file_type",
            "Unknown"
        )


        page = document.metadata.get(
            "page",
            None
        )


        if page is not None:

            page_info = (
                f"Page: {page + 1}"
            )

        else:

            page_info = ""


        formatted.append(
            f"""
Source: {source}
Type: {file_type}
{page_info}

Content:
{document.page_content}
"""
        )


    return "\n\n".join(
        formatted
    )


# =========================================================
# CREATE RAG CHAIN
# =========================================================

def get_chain(
    vectorstore,
    model_name,
    k
):

    retriever = (
        vectorstore.as_retriever(
            search_kwargs={
                "k": k
            }
        )
    )


    llm = ChatOpenAI(
        model=model_name,
        temperature=0
    )


    chain = (

        {
            "context":
                retriever | format_docs,

            "question":
                RunnablePassthrough()
        }

        | RAG_PROMPT

        | llm

        | StrOutputParser()
    )


    return chain


# =========================================================
# FILE UPLOADER
# =========================================================

uploaded_file = st.file_uploader(
    "📂 Upload your file",
    type=SUPPORTED_FILES
)


# =========================================================
# PROCESS FILE
# =========================================================

if uploaded_file:

    file_signature = (
        uploaded_file.name,
        uploaded_file.size,
        chunk_size,
        chunk_overlap,
        model_name
    )


    # -----------------------------------------
    # Rebuild only when needed
    # -----------------------------------------

    if (
        st.session_state.file_sig
        != file_signature
    ):

        with st.spinner(
            "🔄 Processing your file..."
        ):

            try:

                st.session_state.vectorstore = (
                    build_vectorstore(
                        uploaded_file,
                        model_name,
                        chunk_size,
                        chunk_overlap
                    )
                )


                st.session_state.file_sig = (
                    file_signature
                )


                st.session_state.messages = []


                st.success(
                    "✅ File processed successfully!"
                )


            except Exception as e:

                st.error(
                    "❌ Error while processing file:"
                )

                st.exception(e)

                st.stop()


# =========================================================
# CHAT HISTORY
# =========================================================

for message in (
    st.session_state.messages
):

    with st.chat_message(
        message["role"]
    ):

        st.markdown(
            message["content"]
        )


# =========================================================
# CHAT INPUT
# =========================================================

if st.session_state.vectorstore:

    question = st.chat_input(
        "💬 Ask a question about your file..."
    )


    if question:

        # -------------------------------------
        # User message
        # -------------------------------------

        st.session_state.messages.append(
            {
                "role": "user",
                "content": question
            }
        )


        with st.chat_message("user"):

            st.markdown(
                question
            )


        # -------------------------------------
        # Assistant
        # -------------------------------------

        with st.chat_message(
            "assistant"
        ):

            with st.spinner(
                "🤔 Searching the document..."
            ):

                try:

                    # ---------------------------------
                    # Create chain
                    # ---------------------------------

                    chain = get_chain(
                        st.session_state.vectorstore,
                        model_name,
                        top_k
                    )


                    # ---------------------------------
                    # Generate answer
                    # ---------------------------------

                    answer = chain.invoke(
                        question
                    )


                    st.markdown(
                        answer
                    )


                    # ---------------------------------
                    # Save answer
                    # ---------------------------------

                    st.session_state.messages.append(
                        {
                            "role": "assistant",
                            "content": answer
                        }
                    )


                    # ---------------------------------
                    # Retrieve source documents
                    # ---------------------------------

                    retriever = (
                        st.session_state
                        .vectorstore
                        .as_retriever(
                            search_kwargs={
                                "k": top_k
                            }
                        )
                    )


                    source_docs = (
                        retriever.invoke(
                            question
                        )
                    )


                    # ---------------------------------
                    # Show sources
                    # ---------------------------------

                    if source_docs:

                        with st.expander(
                            "📚 Retrieved Sources"
                        ):

                            for i, doc in enumerate(
                                source_docs,
                                start=1
                            ):

                                source = (
                                    doc.metadata.get(
                                        "original_file",
                                        doc.metadata.get(
                                            "source",
                                            "Unknown"
                                        )
                                    )
                                )


                                file_type = (
                                    doc.metadata.get(
                                        "file_type",
                                        "Unknown"
                                    )
                                )


                                page = (
                                    doc.metadata.get(
                                        "page",
                                        None
                                    )
                                )


                                st.write(
                                    f"### Source {i}"
                                )


                                st.write(
                                    f"📄 File: {source}"
                                )


                                st.write(
                                    f"📌 Type: {file_type}"
                                )


                                if page is not None:

                                    st.write(
                                        f"📖 Page: {page + 1}"
                                    )


                                st.write(
                                    doc.page_content[:500]
                                )


                                st.divider()


                except Exception as e:

                    st.error(
                        f"❌ Error: {str(e)}"
                    )


# =========================================================
# NO FILE YET
# =========================================================

else:

    st.info(
        """
        👋 **Welcome to Multiformat RAG Assistant!**

        ### How it works

        **1️⃣ Enter API Key**
        → API key is verified first.

        **2️⃣ Upload File**
        → Upload PDF, Word, Excel, CSV,
        TXT, PowerPoint, image or audio.

        **3️⃣ Extract Content**
        → The system extracts text,
        image information or audio transcript.

        **4️⃣ Chunking**
        → Large content is divided into
        smaller chunks.

        **5️⃣ Embeddings**
        → Each chunk is converted into
        numerical vectors.

        **6️⃣ FAISS**
        → Vectors are stored in FAISS.

        **7️⃣ Ask Question**
        → RAG searches the most relevant
        chunks.

        **8️⃣ Generate Answer**
        → AI answers using only the
        retrieved document information.

        ### Supported Formats

        📄 PDF  
        📝 DOCX  
        📊 XLSX / XLS  
        📋 CSV  
        📃 TXT  
        📽️ PPTX  
        🖼️ JPG / JPEG / PNG  
        🎵 MP3 / WAV  

        ❌ Video is not supported.
        """
    )


# =========================================================
# FOOTER
# =========================================================

st.markdown("---")

st.caption(
    "🤖 Multiformat RAG Assistant | "
    "OpenAI + LangChain + FAISS + Streamlit"
)

# python -m streamlit run multimodalRAGagent.py
